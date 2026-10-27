# Copyright (c) 2026, Farbod Siyahpoosh and contributors
"""Voucher streaming export compatibility gate (5.6.21).

Hardens named-field export rows after party_type / party_name were added.
"""

from __future__ import annotations

import csv
import io
import unittest

import frappe
from frappe.utils import flt

from erpnext_extensions.iran_accounting.account_explorer.api import get_voucher_summary
from erpnext_extensions.iran_accounting.account_explorer.export import (
	_voucher_row_to_export_values,
	build_streaming_voucher_csv,
	build_streaming_voucher_xlsx,
	get_export_columns,
)
from erpnext_extensions.iran_accounting.account_explorer.query_spec import (
	AccountExplorerQuerySpec_from_client,
)
from erpnext_extensions.iran_accounting.account_explorer.voucher_gl_lines import (
	VOUCHER_EXPORT_ROW_FIELDS,
	_export_batch,
	iter_voucher_gl_line_export_batches,
)


COMPANY = "اسپاد فارمد دارو"
FY1405 = {
	"company": COMPANY,
	"fiscal_year": "1405",
	"from_date": "2026-03-21",
	"to_date": "2027-03-20",
}
CUSTOMER = "CUS-00003"


def _payload(axis="voucher", party_scope=None, **extra):
	ctx = {
		"view_axis": axis,
		"detail_mode": "summary",
		"page": 1,
		"page_size": 50,
		"party_scope": party_scope or {},
	}
	ctx.update(extra)
	return {"document_scope": dict(FY1405), "analysis_context": ctx}


class TestVoucherExportCompatV5621(unittest.TestCase):
	@classmethod
	def setUpClass(cls):
		frappe.set_user("Administrator")
		cls.customer_name = frappe.db.get_value("Customer", CUSTOMER, "customer_name")

	def test_export_batch_returns_named_dicts(self):
		spec = AccountExplorerQuerySpec_from_client(
			_payload(
				party_scope={
					"unified_identity": 1,
					"identity_members": [["Customer", CUSTOMER]],
				}
			)
		)
		from erpnext_extensions.iran_accounting.account_explorer.dimension_discovery import (
			filter_usable_gl_dimensions,
			get_discovered_dimensions,
		)

		dimensions = filter_usable_gl_dimensions(get_discovered_dimensions())
		batch = _export_batch(spec, dimensions, batch_size=25, after_name=None)
		self.assertTrue(batch)
		self.assertIsInstance(batch[0], dict)
		for field in VOUCHER_EXPORT_ROW_FIELDS:
			self.assertIn(field, batch[0])
		# Keyset cursor uses named gl_entry, not tuple index 0.
		self.assertTrue(batch[0].get("gl_entry"))

	def test_named_row_export_values_align_headers(self):
		columns = get_export_columns(
			AccountExplorerQuerySpec_from_client(_payload())
		)
		fieldnames = [c["fieldname"] for c in columns]
		row = {
			"gl_entry": "ACC-GLE-TEST",
			"posting_date": "2026-04-01",
			"voucher_type": "Sales Invoice",
			"voucher_no": "SI-1",
			"account_code": "1100",
			"account_name": "Receivable",
			"debit": 100,
			"credit": 0,
			"account_currency": "IRR",
			"remarks": "test",
			"party_type": "Customer",
			"party_name": self.customer_name,
			"party": CUSTOMER,
			"cost_center": "",
			"account": "1100 - Receivable - E",
		}
		values = _voucher_row_to_export_values(row, fieldnames, date_format="gregorian")
		self.assertEqual(len(values), len(fieldnames))
		mapped = dict(zip(fieldnames, values, strict=True))
		self.assertEqual(mapped.get("party"), CUSTOMER)
		self.assertEqual(mapped.get("party_name"), self.customer_name)
		self.assertEqual(mapped.get("party_type"), "Customer")
		self.assertEqual(mapped.get("debit"), 100)
		# Party Name must never overwrite Party ID column.
		self.assertNotEqual(mapped.get("party"), mapped.get("party_name"))

	def test_legacy_15_tuple_still_maps(self):
		fieldnames = ["party_type", "party_name", "party", "debit", "credit"]
		legacy = (
			"GLE-1",
			"2026-04-01",
			"Sales Invoice",
			"SI-1",
			"1100",
			"Receivable",
			50,
			0,
			"IRR",
			"r",
			"Customer",
			self.customer_name,
			CUSTOMER,
			"",
			"1100 - Receivable - E",
		)
		values = _voucher_row_to_export_values(legacy, fieldnames, date_format="gregorian")
		self.assertEqual(values[0], "Customer")
		self.assertEqual(values[1], self.customer_name)
		self.assertEqual(values[2], CUSTOMER)

	def test_legacy_wrong_tuple_length_rejected(self):
		with self.assertRaises(Exception):
			_voucher_row_to_export_values((1, 2, 3), ["party"], date_format="gregorian")

	def test_streaming_csv_party_columns_and_totals(self):
		spec = AccountExplorerQuerySpec_from_client(
			_payload(
				party_scope={
					"unified_identity": 1,
					"identity_members": [["Customer", CUSTOMER]],
				}
			)
		)
		columns = get_export_columns(spec)
		content, totals, total_rows = build_streaming_voucher_csv(
			spec, columns, batch_size=500, date_format="gregorian"
		)
		self.assertGreater(total_rows, 0)
		reader = csv.reader(io.StringIO(content))
		headers = next(reader)
		self.assertIn("Party Type", headers)
		self.assertIn("Party Name", headers)
		self.assertIn("Party", headers)
		idx_type = headers.index("Party Type")
		idx_name = headers.index("Party Name")
		idx_party = headers.index("Party")
		idx_debit = headers.index("Debit")
		idx_credit = headers.index("Credit")
		rows = list(reader)
		self.assertEqual(len(rows), total_rows)
		debit_sum = credit_sum = 0.0
		for row in rows:
			if row[idx_party] == CUSTOMER:
				self.assertEqual(row[idx_type], "Customer")
				self.assertEqual(row[idx_name], self.customer_name)
			debit_sum += flt(row[idx_debit])
			credit_sum += flt(row[idx_credit])
		self.assertAlmostEqual(debit_sum, flt(totals.get("debit") or totals.get("scoped_debit")), places=2)
		self.assertAlmostEqual(credit_sum, flt(totals.get("credit") or totals.get("scoped_credit")), places=2)

	def test_streaming_xlsx_opens_and_preserves_party_text(self):
		from openpyxl import load_workbook

		spec = AccountExplorerQuerySpec_from_client(
			_payload(
				party_scope={
					"unified_identity": 1,
					"identity_members": [["Customer", CUSTOMER]],
				}
			)
		)
		columns = get_export_columns(spec)
		content, totals, total_rows = build_streaming_voucher_xlsx(
			spec, columns, batch_size=500, date_format="jalali"
		)
		self.assertGreater(total_rows, 0)
		self.assertTrue(content)
		wb = load_workbook(io.BytesIO(content), read_only=True)
		ws = wb.active
		headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
		self.assertIn("Party Name", headers)
		self.assertIn("Party", headers)
		name_i = headers.index("Party Name")
		party_i = headers.index("Party")
		found = False
		for row in ws.iter_rows(min_row=2, values_only=True):
			if row[party_i] == CUSTOMER:
				self.assertEqual(row[name_i], self.customer_name)
				found = True
				break
		self.assertTrue(found)
		wb.close()

	def test_selected_export_only_selected_keys(self):
		from erpnext_extensions.iran_accounting.account_explorer.export import collect_export_rows

		summary = get_voucher_summary(
			_payload(
				party_scope={
					"unified_identity": 1,
					"identity_members": [["Customer", CUSTOMER]],
				}
			)
		)
		rows = summary.get("rows") or []
		self.assertGreaterEqual(len(rows), 2)
		selected = [rows[0]["row_key"], rows[1]["row_key"]]
		spec = AccountExplorerQuerySpec_from_client(
			_payload(
				party_scope={
					"unified_identity": 1,
					"identity_members": [["Customer", CUSTOMER]],
				}
			)
		)
		exported, _totals, total_rows = collect_export_rows(
			spec, row_scope="selected", selected_row_keys=selected
		)
		self.assertEqual(total_rows, 2)
		self.assertEqual(len(exported), 2)
		exported_keys = {r.get("row_key") for r in exported}
		self.assertEqual(exported_keys, set(selected))

	def test_ui_summary_matches_export_debit_for_customer(self):
		summary = get_voucher_summary(
			_payload(
				party_scope={
					"unified_identity": 1,
					"identity_members": [["Customer", CUSTOMER]],
				}
			)
		)
		spec = AccountExplorerQuerySpec_from_client(
			_payload(
				party_scope={
					"unified_identity": 1,
					"identity_members": [["Customer", CUSTOMER]],
				}
			)
		)
		# Independent SQL
		sql = frappe.db.sql(
			"""
			SELECT COALESCE(SUM(debit),0), COALESCE(SUM(credit),0), COUNT(*)
			FROM `tabGL Entry`
			WHERE company=%s AND is_cancelled=0
				AND posting_date BETWEEN %s AND %s
				AND party_type=%s AND party=%s
			""",
			(COMPANY, FY1405["from_date"], FY1405["to_date"], "Customer", CUSTOMER),
		)[0]
		# Export stream totals for same identity scope
		columns = get_export_columns(spec)
		_content, totals, total_rows = build_streaming_voucher_csv(
			spec, columns, batch_size=2000, date_format="gregorian"
		)
		# Export may include only non-zero rows depending on hide_zero_rows;
		# compare UI totals to export totals for same spec.
		self.assertAlmostEqual(
			flt(totals.get("debit") or totals.get("scoped_debit")),
			flt((summary.get("totals") or {}).get("debit") or (summary.get("totals") or {}).get("scoped_debit")),
			places=2,
		)
		self.assertGreater(total_rows, 0)
		# Soft bound vs independent SQL (UI may exclude opening/PCV per status filters).
		self.assertLessEqual(flt(totals.get("debit")), flt(sql[0]) + 1)
		_ = sql

	def test_iterator_yields_dicts_not_positional(self):
		spec = AccountExplorerQuerySpec_from_client(
			_payload(
				party_scope={
					"unified_identity": 1,
					"identity_members": [["Customer", CUSTOMER]],
				}
			)
		)
		seen = False
		for batch, _totals, total_rows in iter_voucher_gl_line_export_batches(spec, batch_size=50):
			if not batch:
				continue
			seen = True
			self.assertIsInstance(batch[0], dict)
			self.assertIn("party_name", batch[0])
			self.assertIn("party", batch[0])
			self.assertGreater(total_rows, 0)
			break
		self.assertTrue(seen)
